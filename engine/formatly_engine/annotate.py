"""Writing the issues into the document as Word comments."""

from __future__ import annotations

from collections import Counter, defaultdict

from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from .checks.base import Context
from .issues import Issue

LABEL = {"error": "Must fix", "warning": "Check", "suggestion": "Suggestion"}


def _anchor_paragraph(ctx: Context, block_index: int | None) -> Paragraph | None:
    blocks = ctx.blocks
    if block_index is None:
        block_index = next((b.index for b in blocks if b.kind == "p" and not b.is_empty), None)
        if block_index is None:
            return None
    b = blocks[block_index]
    p_el = b.el if b.kind == "p" else next(b.el.iter(qn("w:p")), None)
    if p_el is None:
        return None
    return Paragraph(p_el, ctx.model.document._body)


def _runs_for(paragraph: Paragraph):
    runs = paragraph.runs
    if not runs:
        runs = [paragraph.add_run("")]
    return runs


def _line(issue: Issue) -> str:
    text = f"{LABEL[issue.severity]}: {issue.message}"
    if issue.hint:
        text += f" How to fix: {issue.hint}"
    if issue.fixable:
        text += " [Auto-fix can correct this.]"
    return text


def annotate(ctx: Context, report: dict) -> int:
    """Add comments to ctx.model.document in place. Returns the number of comments added."""
    rules = ctx.rules.report
    doc = ctx.model.document
    per_rule_limit = max(1, rules.comments_per_rule)

    total_by_rule = Counter(i.rule for i in ctx.log.issues)
    shown_by_rule: Counter = Counter()
    by_block: dict[int | None, list[Issue]] = defaultdict(list)
    for issue in ctx.log.issues:
        if shown_by_rule[issue.rule] >= per_rule_limit:
            continue
        shown_by_rule[issue.rule] += 1
        by_block[issue.block].append(issue)

    # Note on the last shown instance of a rule that has more instances than shown.
    last_shown: dict[str, Issue] = {}
    for issues in by_block.values():
        for i in issues:
            last_shown[i.rule] = i
    overflow = {id(last_shown[r]): total_by_rule[r] - shown_by_rule[r]
                for r in shown_by_rule if total_by_rule[r] > shown_by_rule[r]}

    added = 0
    summary = _anchor_paragraph(ctx, None)
    if summary is not None:
        c = report["counts"]
        text = (f"Formatly check: score {report['score']}/100. {c['errors']} to fix, {c['warnings']} to check, "
                f"{c['suggestions']} suggestions. Rulebook: {report['rulebook']['name']}.")
        comment = doc.add_comment(_runs_for(summary), text=text, author=rules.comment_author,
                                  initials=rules.comment_initials)
        if c["fixable"]:
            comment.add_paragraph(f"{c['fixable']} of these can be corrected automatically with Auto-fix.")
        added += 1

    for block, issues in by_block.items():
        paragraph = _anchor_paragraph(ctx, block)
        if paragraph is None:
            continue
        order = {"error": 0, "warning": 1, "suggestion": 2}
        issues = sorted(issues, key=lambda i: order[i.severity])
        comment = doc.add_comment(_runs_for(paragraph), text=_line(issues[0]),
                                  author=rules.comment_author, initials=rules.comment_initials)
        for i in issues[1:]:
            comment.add_paragraph(_line(i))
        for i in issues:
            more = overflow.get(id(i))
            if more:
                comment.add_paragraph(f"The same problem appears in {more} more place{'s' if more > 1 else ''}; "
                                      "see the full report.")
        added += 1
    return added
