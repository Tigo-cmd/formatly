import pytest

from formatly_engine.checks.references import parse_citations
from formatly_engine.docmodel import Block
from formatly_engine.rulebook import LineSpacing, load_rulebook, rulebook_from_dict
from formatly_engine.units import parse_length


@pytest.mark.parametrize("text,twips", [("1in", 1440), ("2.54cm", 1440), ("25.4mm", 1440), ("72pt", 1440),
                                        (12, 240), ("1.5 in", 2160)])
def test_parse_length(text, twips):
    assert parse_length(text) == twips


def test_parse_length_rejects_nonsense():
    with pytest.raises(ValueError):
        parse_length("wide")


@pytest.mark.parametrize("value,rule,amount", [(2, "multiple", 2.0), ("double", "multiple", 2.0),
                                               ("1.5", "multiple", 1.5), ("exactly 12pt", "exact", 12.0)])
def test_line_spacing(value, rule, amount):
    ls = LineSpacing.model_validate(LineSpacing.parse(value))
    assert (ls.rule, ls.value) == (rule, amount)


def test_generic_rulebook_loads():
    rb = load_rulebook("generic")
    assert rb.body.font == "Times New Roman"
    assert rb.page.margins.left == 2160
    assert [c.number for c in rb.structure.chapters] == [1, 2, 3, 4, 5]


def test_extending_overrides_only_what_is_given():
    rb = rulebook_from_dict({"name": "Dept X", "extends": "default",
                             "body": {"line_spacing": 1.5}, "page": {"margins": {"left": "1in", "right": "1in",
                                                                                   "top": "1in", "bottom": "1in"}}})
    assert rb.name == "Dept X"
    assert rb.body.line_spacing.value == 1.5
    assert rb.body.font == "Times New Roman"  # inherited
    assert rb.page.margins.left == 1440


def test_unknown_field_is_an_error():
    with pytest.raises(Exception):
        rulebook_from_dict({"name": "x", "body": {"colour": "red"}})


def _block(text):
    return Block(index=0, kind="p", el=None, section=0, text=text)


def test_citation_parsing():
    text = ("Mobile money raises sales (Okafor & Bello, 2021; see Eze et al., 2020, p. 4). "
            "According to Adeyemi (2019) and Musa and Ali (2018), results vary. In Nigeria (2019) adoption grew.")
    found = {(c.surname, c.year) for c in parse_citations([_block(text)])}
    assert found == {("okafor", "2021"), ("eze", "2020"), ("adeyemi", "2019"), ("musa", "2018")}


def test_citation_with_several_years_and_no_date():
    found = {(c.surname, c.year) for c in parse_citations([_block("(Obi, 2019, 2021; Federal Ministry of Health, n.d.)")])}
    assert found == {("obi", "2019"), ("obi", "2021"), ("federal", "n.d.")}


def test_sentence_about_a_chapter_is_not_a_chapter_heading():
    import io

    import docx

    from formatly_engine import check
    from samples import build_project

    d = docx.Document(io.BytesIO(build_project()))
    target = next(p for p in d.paragraphs if p.text.startswith("This study examines"))
    target.insert_paragraph_before("Chapter two reviews related literature.")
    buf = io.BytesIO()
    d.save(buf)
    result = check(buf.getvalue(), "generic")
    assert [c["number"] for c in result.report["document"]["chapters"]] == [1, 2, 3, 4, 5]
    assert not any(i.rule.startswith("structure.") for i in result.issues)


def test_tab_stop_settings_are_not_text():
    """WPS and Word store tab stops as <w:tab> inside paragraph settings."""
    import docx
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Inches

    from formatly_engine.docmodel import paragraph_text

    d = docx.Document()
    p = d.add_paragraph("Declaration")
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6), WD_TAB_ALIGNMENT.RIGHT)
    assert paragraph_text(p._p) == "Declaration"


def test_long_numbered_heading_is_a_heading():
    import io

    import docx

    from formatly_engine import check
    from samples import build_standard_project

    d = docx.Document(io.BytesIO(build_standard_project()))
    target = [p for p in d.paragraphs if p.text.startswith("2.3 ")][-1]  # the body heading, not the contents line
    long_title = ("2.2.1 CONTROLLING A DRONE WITH HAND GESTURES USING MEDIAPIPE HAND KEYPOINT DETECTION "
                  "AND A SIMPLE NEURAL NETWORK TO RECOGNISE GESTURES")
    h = target.insert_paragraph_before(long_title)
    h.runs[0].bold = True
    buf = io.BytesIO()
    d.save(buf)
    result = check(buf.getvalue())
    levels = {info.text: info.level for info in result.context.outline.headings.values()}
    assert levels.get(long_title) == 3


def test_must_fix_problems_keep_the_score_out_of_the_good_range():
    from formatly_engine import check
    from samples import build_standard_project

    one_problem = check(build_standard_project(missing_scope=True)).report
    assert one_problem["counts"]["errors"] == 1 and one_problem["score"] <= 84
    only_suggestion = check(build_standard_project(old_reference=True)).report
    assert only_suggestion["score"] == 100  # suggestions do not cost points
