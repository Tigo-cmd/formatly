import io

import docx
import pytest

from formatly_engine import check as _check
from formatly_engine import fix as _fix
from formatly_engine import load_rulebook
from samples import build_messy_project, build_project

# The samples follow the generic rulebook (A4, 1.5 in binding margin, APA 7).
GENERIC = load_rulebook("generic")


def check(data, rules=None):
    return _check(data, rules or GENERIC)


def fix(data, rules=None):
    return _fix(data, rules or GENERIC)

# One typical mistake -> the rule that must catch it, and whether auto-fix must remove it.
MISTAKES = [
    ({"body_font": "Calibri"}, "body.font", True),
    ({"body_size": 11}, "body.size", True),
    ({"spacing": 1.5}, "body.line_spacing", True),
    ({"justify": False}, "body.alignment", True),
    ({"letter_paper": True}, "page.size", True),
    ({"even_margins": True}, "page.margins", True),
    ({"chapter_digits": True}, "headings.chapter_label", True),
    ({"headings_not_bold": True}, "headings.format", True),
    ({"theme_heading_font": True}, "headings.format", True),
    ({"chapter_same_page": True}, "headings.new_page", True),
    ({"missing_dedication": True}, "structure.missing", True),
    ({"abstract_after_toc": True}, "structure.order", True),
    ({"table_caption_below": True}, "captions.position", True),
    ({"unsorted_references": True}, "references.order", True),
    ({"no_hanging_indent": True}, "references.hanging_indent", True),
    ({"highlight": True}, "hygiene.highlight", True),
    ({"double_spaces": True}, "hygiene.double_space", True),
    ({"no_page_numbers": True}, "numbering.present", True),
    ({"no_section_break": True}, "numbering.section_break", True),
    ({"uncited_reference": True}, "references.uncited", False),
    ({"missing_reference": True}, "references.citation_missing", False),
    ({"typed_toc": True}, "toc.automatic", False),
    ({"long_abstract": True}, "limits.abstract", False),
    ({"wrong_subheading_number": True}, "headings.numbering", False),
]


def rules_of(result):
    return {i.rule for i in result.issues}


def test_correct_document_passes():
    result = check(build_project())
    assert [i.message for i in result.issues] == []
    assert result.report["score"] == 100
    found = [s["id"] for s in result.report["document"]["sections_found"]]
    assert found == ["certification", "dedication", "acknowledgements", "abstract", "table_of_contents",
                     "chapter_1", "chapter_2", "chapter_3", "chapter_4", "chapter_5", "references"]
    assert result.report["document"]["captions"] == 2


@pytest.mark.parametrize("options,rule,fixable", MISTAKES, ids=[m[1] + ":" + next(iter(m[0])) for m in MISTAKES])
def test_mistake_is_found(options, rule, fixable):
    result = check(build_project(**options))
    assert rule in rules_of(result), [i.message for i in result.issues]
    assert result.report["score"] < 100
    # The mistake must not cause unrelated errors elsewhere.
    unrelated = {i.rule for i in result.issues if i.severity == "error"} - {rule}
    looks = {"references.font", "references.size", "references.line_spacing", "captions.format",
             "captions.table_text"}  # the body font and size also apply to these
    allowed = {
        "body_font": looks, "body_size": looks, "spacing": looks,
        "missing_dedication": set(), "abstract_after_toc": set(),
        "no_section_break": {"numbering.preliminaries", "numbering.main"},
        "theme_heading_font": set(), "long_abstract": set(),
    }
    assert unrelated <= allowed.get(next(iter(options)), set()), unrelated


@pytest.mark.parametrize("options,rule,fixable", [m for m in MISTAKES if m[2]],
                         ids=[m[1] + ":" + next(iter(m[0])) for m in MISTAKES if m[2]])
def test_mistake_is_fixed(options, rule, fixable):
    result = fix(build_project(**options))
    assert rule in result.applied
    remaining = rules_of(result.after) - {"hygiene.placeholder"}
    assert rule not in remaining, [i.message for i in result.after.issues]
    assert remaining == set(), [i.message for i in result.after.issues]


def test_fixing_a_correct_document_changes_nothing():
    result = fix(build_project())
    assert result.applied == {}
    assert result.after.report["score"] == 100


def test_messy_document_end_to_end():
    result = fix(build_messy_project())
    assert result.before["score"] < 50
    assert result.after.report["score"] >= result.before["score"] + 30
    assert result.after.report["score"] <= 84  # the abstract is still too long: must-fix
    left = rules_of(result.after)
    # What remains is only what the writer must do themselves.
    assert left <= {"references.citation_missing", "references.uncited", "limits.abstract",
                    "hygiene.placeholder"}, left
    assert not any(i.fixable for i in result.after.issues)
    docx.Document(io.BytesIO(result.fixed))  # still a valid document


def test_annotated_copy_has_comments():
    result = check(build_messy_project())
    data = result.annotated_bytes()
    comments = list(docx.Document(io.BytesIO(data)).comments)
    assert len(comments) > 10
    assert comments[0].text.startswith("Formatly check: score")
    assert any("Required: Times New Roman" in c.text for c in comments)


def test_comment_limit_per_rule():
    rb = load_rulebook("generic")
    rb.report.comments_per_rule = 2
    result = check(build_project(body_font="Calibri"), rb)
    data = result.annotated_bytes()
    comments = list(docx.Document(io.BytesIO(data)).comments)
    font_comments = [c for c in comments if "Must fix: Font is Calibri" in c.text]
    assert len(font_comments) == 2
    assert any("more place" in c.text for c in font_comments)


def test_style_hint_when_normal_style_is_wrong():
    result = check(build_project(body_font="Calibri"))
    font_issue = next(i for i in result.issues if i.rule == "body.font")
    assert "Normal" in font_issue.hint


def test_documents_written_by_libreoffice(tmp_path):
    """Real files come from many programs. LibreOffice writes Word XML differently
    (e.g. section breaks inside the table of contents box); results must not change."""
    import shutil
    import subprocess

    soffice = shutil.which("soffice")
    if not soffice:
        pytest.skip("LibreOffice not installed")
    for name, data, expect_clean in (("good", build_project(), True), ("messy", build_messy_project(), False)):
        src = tmp_path / f"{name}-src.docx"
        src.write_bytes(data)
        subprocess.run([soffice, "--headless", "--convert-to", "odt", "--outdir", str(tmp_path), str(src)],
                       check=True, capture_output=True, timeout=180)
        odt = tmp_path / f"{name}-src.odt"
        subprocess.run([soffice, "--headless", "--convert-to", "docx:MS Word 2007 XML", "--outdir",
                        str(tmp_path / "rt"), str(odt)], check=True, capture_output=True, timeout=180)
        converted = (tmp_path / "rt" / f"{name}-src.docx").read_bytes()
        result = check(converted)
        if expect_clean:
            assert [i.message for i in result.issues] == []
        else:
            fixed = fix(converted)
            assert not any(i.fixable for i in fixed.after.issues)
            assert fixed.after.report["score"] >= fixed.before["score"] + 30
