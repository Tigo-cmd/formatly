"""The Formatly Standard rulebooks (default, seminar, postgraduate), built from
the team's template documents and written guidance."""

import io

import docx
import pytest

from formatly_engine import check, fix
from samples import build_project, build_standard_project


def rules_of(result):
    return {i.rule for i in result.issues}


@pytest.mark.parametrize("variant,rulebook", [({}, "default"), ({"seminar": True}, "seminar"),
                                              ({"a4": True}, "default"),
                                              ({"report_kind": "DISSERTATION"}, "postgraduate")])
def test_correct_standard_document_passes(variant, rulebook):
    result = check(build_standard_project(**variant), rulebook)
    assert [i.message for i in result.issues] == []
    assert result.report["score"] == 100


# mistake -> rule that must catch it, severity, auto-fixable
MISTAKES = [
    ("certification_supervisor", "structure.content", "error", False),
    ("missing_scope", "structure.subsections", "error", False),
    ("old_reference", "references.age", "suggestion", False),
    ("abstract_two_paragraphs", "structure.content", "error", False),
    ("abstract_double", "body.line_spacing", "error", True),
    ("toc_wrong_title", "toc.up_to_date", "warning", False),
    ("toc_without_numbers", "toc.up_to_date", "warning", False),
    ("list_title_mismatch", "toc.lists", "warning", False),
    ("no_list_of_tables", "structure.missing", "error", True),
    ("blank_lines_before_chapter", "headings.new_page", "error", True),
    ("split_reference", "references.split", "error", True),
    ("reference_tabs", "references.layout", "error", True),
    ("title_page_no_reg_number", "structure.title_page_content", "error", False),
]


@pytest.mark.parametrize("flag,rule,severity,fixable", MISTAKES, ids=[m[0] for m in MISTAKES])
def test_standard_mistake_is_found(flag, rule, severity, fixable):
    result = check(build_standard_project(**{flag: True}))
    found = [i for i in result.issues if i.rule == rule]
    assert found, [i.message for i in result.issues]
    assert {i.severity for i in found} == {severity}
    assert all(i.fixable == fixable for i in found)
    others = rules_of(result) - {rule}
    # Leaving out the list of tables also leaves its table out of every list.
    allowed = {"no_list_of_tables": {"toc.lists"},
               "split_reference": {"references.layout"}}  # the wrapped line starts with a tab
    assert others <= allowed.get(flag, set()), [i.message for i in result.issues if i.rule != rule]


@pytest.mark.parametrize("flag,rule", [(m[0], m[1]) for m in MISTAKES if m[3]], ids=[m[0] for m in MISTAKES if m[3]])
def test_standard_mistake_is_fixed(flag, rule):
    result = fix(build_standard_project(**{flag: True}))
    remaining = rules_of(result.after)
    assert rule not in remaining, [i.message for i in result.after.issues]
    assert not any(i.fixable for i in result.after.issues)


def test_certification_wording_points_at_the_line():
    result = check(build_standard_project(certification_supervisor=True))
    msgs = sorted(i.message for i in result.issues)
    assert msgs == ["Certification must include “External Examiner”.",
                    "Certification should not say “External Supervisor”."]
    forbidden = next(i for i in result.issues if "should not" in i.message)
    assert result.context.blocks[forbidden.block].text == "External Supervisor"


def test_split_reference_is_joined_without_changing_words():
    before = build_standard_project(split_reference=True)
    result = fix(before)
    doc = docx.Document(io.BytesIO(result.fixed))
    refs = [p.text for p in doc.paragraphs if p.text.startswith("Okafor")]
    assert len(refs) == 1
    assert refs[0].endswith("Automatic watering with solar pumps. Lagos: University Press.")
    assert "\t" not in refs[0] and "  " not in refs[0]


def test_blank_lines_are_replaced_by_a_page_break():
    result = check(build_standard_project(blank_lines_before_chapter=True))
    issue = next(i for i in result.issues if i.rule == "headings.new_page")
    assert issue.data["blank_lines"] == 8
    assert "empty lines" in issue.message
    fixed = fix(build_standard_project(blank_lines_before_chapter=True))
    doc = docx.Document(io.BytesIO(fixed.fixed))
    texts = [p.text for p in doc.paragraphs]
    i = texts.index("CHAPTER TWO")
    assert texts[i - 1] != ""  # the empty lines are gone


def test_postgraduate_treats_old_sources_as_errors():
    result = check(build_standard_project(old_reference=True, report_kind="THESIS"), "postgraduate")
    ages = [i for i in result.issues if i.rule == "references.age"]
    assert ages and all(i.severity == "error" for i in ages)


def test_seminar_needs_only_two_chapters_and_background_of_the_study():
    result = check(build_standard_project(seminar=True), "seminar")
    assert result.issues == []
    # A seminar title page must say "seminar".
    project = check(build_standard_project(), "seminar")
    assert "structure.title_page_content" in rules_of(project)


def test_letter_and_a4_are_both_accepted_and_fix_keeps_the_paper():
    result = fix(build_project(even_margins=True), "default")
    doc = docx.Document(io.BytesIO(result.fixed))
    assert round(doc.sections[0].page_width.mm) == 210  # A4 stays A4


def test_typed_toc_headings_are_not_taken_for_chapters():
    result = check(build_standard_project())
    assert [h["number"] for h in result.report["document"]["chapters"]] == [1, 2, 3, 4, 5]
    toc = result.context.outline.hit("table_of_contents")
    chapter_one = result.context.outline.hit("chapter_1")
    assert chapter_one.block > toc.block + 10  # the real chapter, not the contents line
